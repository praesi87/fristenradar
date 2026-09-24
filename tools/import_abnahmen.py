#!/usr/bin/env python3
"""import_abnahmen.py — befüllt das Gewährleistungsregister aus eigenen Dokumenten.

Liest Abnahmeprotokolle, Bauverträge und Bürgschaftsurkunden (PDF/DOCX) aus einem
Ordner und erzeugt daraus Registerzeilen im Schema von gewaehrleistungsregister.csv.

Grundregel: Es wird NICHTS geraten. Jedes Feld wird entweder wörtlich aus dem
Dokument belegt (mit Fundstelle) oder bleibt leer und wird als Prüfbedarf markiert.
Ein leeres Feld ist ein Arbeitsauftrag, kein Fehler — ein erfundenes Feld wäre ein
Haftungsrisiko (falscher Fristbeginn = verjährter Anspruch).

Aufruf:
    python3 import_abnahmen.py <eingangsordner> [-o rohimport.csv]
    python3 import_abnahmen.py --selftest        # erzeugt Testdokumente und prüft sich selbst
"""
import argparse, csv, datetime as dt, pathlib, re, sys, unicodedata

SPALTEN = ["Projekt", "Adresse", "PLZ", "Ort", "Region", "Leistung/Gewerk",
           "Auftragnehmer", "Rolle AN", "Vertragsart", "Abnahmedatum", "Fristende",
           "Bauherr", "Ansprechpartner Bauherr", "Property Manager", "Sicherheit",
           "Status", "Bemerkung"]
ZUSATZ = ["_Quelldatei", "_Fundstellen", "_Pruefbedarf"]

# Orte/Kreise, die den Verflechtungsraum Rhein-Main abdecken (Region-Vorbelegung).
# Bewusst konservativ: kein Treffer -> Region bleibt leer und wird zum Prüfbedarf.
RHEIN_MAIN = {
    "frankfurt am main", "frankfurt", "offenbach am main", "offenbach", "wiesbaden",
    "mainz", "darmstadt", "hanau", "rüsselsheim am main", "rüsselsheim", "bad homburg",
    "bad homburg vor der höhe", "oberursel", "oberursel (taunus)", "friedberg",
    "friedberg (hessen)", "bad vilbel", "maintal", "dreieich", "neu-isenburg",
    "langen", "langen (hessen)", "mörfelden-walldorf", "kelkheim", "kelkheim (taunus)",
    "hofheim", "hofheim am taunus", "eschborn", "schwalbach am taunus", "kronberg",
    "kronberg im taunus", "königstein", "königstein im taunus", "bad soden",
    "bad soden am taunus", "flörsheim", "flörsheim am main", "hattersheim",
    "hattersheim am main", "ginsheim-gustavsburg", "raunheim", "kelsterbach",
    "rodgau", "obertshausen", "heusenstamm", "dietzenbach", "rödermark", "seligenstadt",
    "groß-gerau", "gross-gerau", "bischofsheim", "nauheim", "trebur", "griesheim",
    "weiterstadt", "pfungstadt", "bad nauheim", "butzbach", "karben", "nidderau",
    "bruchköbel", "gelnhausen", "aschaffenburg", "bingen am rhein", "ingelheim am rhein",
    "budenheim", "hochheim am main", "eltville am rhein", "taunusstein", "idstein",
    "bad schwalbach", "limburg an der lahn", "usingen", "friedrichsdorf", "steinbach",
    "sulzbach", "liederbach", "eppstein", "kriftel", "münster", "rüdesheim am rhein",
}

MONATE = {"januar": 1, "februar": 2, "märz": 3, "maerz": 3, "april": 4, "mai": 5,
          "juni": 6, "juli": 7, "august": 8, "september": 9, "oktober": 10,
          "november": 11, "dezember": 12}

# --------------------------------------------------------------------------- Text

def text_aus_pdf(p):
    import pdfplumber
    seiten = []
    with pdfplumber.open(p) as pdf:
        for i, s in enumerate(pdf.pages, 1):
            seiten.append((i, s.extract_text() or ""))
    return seiten

def text_aus_docx(p):
    import docx
    d = docx.Document(str(p))
    teile = [par.text for par in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            teile.append("\t".join(c.text for c in row.cells))
    return [(1, "\n".join(teile))]

def dokument_lesen(p):
    suf = p.suffix.lower()
    if suf == ".pdf":
        return text_aus_pdf(p)
    if suf in (".docx", ".dotx"):
        return text_aus_docx(p)
    if suf in (".txt", ".md"):
        return [(1, p.read_text(encoding="utf-8", errors="replace"))]
    raise ValueError(f"nicht unterstützt: {suf}")

def norm(s):
    s = unicodedata.normalize("NFKC", s or "")
    return re.sub(r"[ \t ]+", " ", s)

# ------------------------------------------------------------------------ Parsing

def datum_parsen(roh):
    roh = roh.strip().rstrip(".,;")
    m = re.fullmatch(r"(\d{1,2})\.\s*(\d{1,2})\.\s*(\d{4})", roh)
    if m:
        d, mo, y = (int(x) for x in m.groups())
    else:
        m = re.fullmatch(r"(\d{1,2})\.\s*([A-Za-zÄÖÜäöüß]+)\s+(\d{4})", roh)
        if m and m.group(2).lower() in MONATE:
            d, mo, y = int(m.group(1)), MONATE[m.group(2).lower()], int(m.group(3))
        else:
            m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", roh)
            if not m:
                return None
            y, mo, d = (int(x) for x in m.groups())
    try:
        datum = dt.date(y, mo, d)
    except ValueError:
        return None
    # Plausibilitätsgrenzen: Bauabnahmen vor 1990 oder in ferner Zukunft sind Lesefehler.
    if not (dt.date(1990, 1, 1) <= datum <= dt.date.today() + dt.timedelta(days=365)):
        return None
    return datum

DATUM = r"(\d{1,2}\.\s*\d{1,2}\.\s*\d{4}|\d{1,2}\.\s*[A-Za-zÄÖÜäöüß]+\s+\d{4}|\d{4}-\d{2}-\d{2})"

# Reihenfolge = Priorität. Nur Muster, die das Datum eindeutig der ABNAHME zuordnen.
ABNAHME_MUSTER = [
    # {0,2} Zwischenwörter fangen Varianten wie "förmlichen", "foermlichen",
    # "rechtsgeschäftlichen", "gemeinsamen" ab, ohne das Muster zu entgrenzen.
    (r"(?:Tag|Datum|Zeitpunkt)\s+der\s+(?:\w+\s+){0,2}Abnahme\s*[:\-]?\s*" + DATUM, "Tag/Datum der Abnahme"),
    (r"Abnahmedatum\s*[:\-]?\s*" + DATUM, "Abnahmedatum"),
    (r"(?:\w+\s+){0,1}Abnahme\s+(?:erfolgte|erfolgt)?\s*am\s+" + DATUM, "Abnahme am"),
    (r"Die\s+Abnahme\s+(?:fand|wurde)[^.\n]{0,60}?am\s+" + DATUM, "Die Abnahme … am"),
    (r"Abnahmebegehung\s+(?:vom|am)\s+" + DATUM, "Abnahmebegehung"),
    (r"abgenommen\s+am\s+" + DATUM, "abgenommen am"),
]

FELD_MUSTER = {
    "Projekt": [r"(?:Bauvorhaben|Projekt|Objekt|BV)\s*[:\-]\s*(.+)"],
    "Auftragnehmer": [r"(?:Auftragnehmer|AN|Unternehmer|ausführende[s]?\s+Unternehmen)\s*[:\-]\s*(.+)"],
    "Bauherr": [r"(?:Auftraggeber|AG|Bauherr(?:in)?)\s*[:\-]\s*(.+)"],
    "Leistung/Gewerk": [r"(?:Leistung|Gewerk|Leistungsumfang|Bauleistung)\s*[:\-]\s*(.+)"],
    "Sicherheit": [r"(?:Sicherheit|Gewährleistungsbürgschaft|Sicherheitseinbehalt|Bürgschaft)\s*[:\-]\s*(.+)"],
}

def erstes_treffer_feld(text, muster_liste, maxlen=160):
    for mu in muster_liste:
        m = re.search(mu, text, re.IGNORECASE)
        if m:
            wert = norm(m.group(1)).strip(" .;,")
            wert = re.split(r"\s{3,}|\t", wert)[0].strip()
            if 1 < len(wert) <= maxlen:
                return wert, m.group(0)[:90]
    return "", ""

def vertragsart_bestimmen(text):
    """VOB/B nur bei belegter Einbeziehung. Sonst BGB — mit Prüfvermerk."""
    t = text.lower()
    vob_treffer = re.search(r"vob\s*/\s*b|vergabe-?\s*und\s+vertragsordnung", t)
    if vob_treffer:
        # Gegenprobe: explizite Abbedingung
        if re.search(r"vob\s*/\s*b\s+(?:findet|ist)\s+(?:keine|nicht)", t):
            return "BGB", "VOB/B ausdrücklich nicht einbezogen"
        return "VOB/B", "VOB/B im Dokument benannt"
    if re.search(r"\bbgb\b|§\s*63[045]|werkvertrag", t):
        return "BGB", "BGB/Werkvertrag benannt"
    return "", ""

def adresse_finden(text):
    m = re.search(r"([A-ZÄÖÜ][\wÄÖÜäöüß\.\- ]{2,40}(?:straße|strasse|str\.|weg|allee|platz|gasse|ring|damm))\s+(\d{1,4}\s*[a-zA-Z]?)\s*[,\n]\s*(\d{5})\s+([A-ZÄÖÜ][\wÄÖÜäöüß\.\- ]+)", text)
    if m:
        return f"{norm(m.group(1)).strip()} {m.group(2).strip()}", m.group(3), norm(m.group(4)).strip(" .,;"), m.group(0)[:90]
    m = re.search(r"\b(\d{5})\s+([A-ZÄÖÜ][\wÄÖÜäöüß\.\- ]{2,40})", text)
    if m:
        return "", m.group(1), norm(m.group(2)).strip(" .,;"), m.group(0)[:60]
    return "", "", "", ""

def region_bestimmen(ort):
    o = norm(ort).lower().strip()
    o = re.sub(r"\s*\(.*?\)\s*$", "", o).strip()
    return "Rhein-Main" if o in RHEIN_MAIN else ""

# ---------------------------------------------------------------------- Extraktion

def dokument_auswerten(pfad):
    seiten = dokument_lesen(pfad)
    volltext = norm("\n".join(t for _, t in seiten))
    fund, pruef = [], []
    zeile = {k: "" for k in SPALTEN}

    # 1) Abnahmedatum — das kritischste Feld, daher strengste Regeln.
    abnahme, treffer_label = None, ""
    for mu, label in ABNAHME_MUSTER:
        for m in re.finditer(mu, volltext, re.IGNORECASE):
            d = datum_parsen(m.group(1))
            if d:
                abnahme, treffer_label = d, label
                fund.append(f"Abnahmedatum [{label}]: „{norm(m.group(0))[:70]}“")
                break
        if abnahme:
            break
    if abnahme:
        zeile["Abnahmedatum"] = abnahme.isoformat()
        # Warnen, wenn mehrere abweichende Abnahmedaten im Dokument stehen (Teilabnahmen!)
        alle = set()
        for mu, _ in ABNAHME_MUSTER:
            for m in re.finditer(mu, volltext, re.IGNORECASE):
                d = datum_parsen(m.group(1))
                if d:
                    alle.add(d)
        if len(alle) > 1:
            pruef.append(f"MEHRERE Abnahmedaten gefunden ({', '.join(d.isoformat() for d in sorted(alle))}) – Teilabnahmen? Fristbeginn je Teilabnahme separat prüfen")
    else:
        pruef.append("KEIN Abnahmedatum sicher erkannt – manuell eintragen (Fristbeginn!)")

    # 2) Vorbehalte / Teilabnahme / fiktive Abnahme -> immer Prüfbedarf
    if re.search(r"unter\s+Vorbehalt|Vorbehalt\s+der\s+Vertragsstrafe", volltext, re.IGNORECASE):
        pruef.append("Abnahme unter Vorbehalt vermerkt – Vertragsstrafenvorbehalt prüfen")
    if re.search(r"Teilabnahme", volltext, re.IGNORECASE):
        pruef.append("Teilabnahme erwähnt – eigene Fristzeile je Teilabnahme anlegen")
    if re.search(r"Zustandsfeststellung|§\s*650g", volltext, re.IGNORECASE):
        pruef.append("Zustandsfeststellung § 650g erwähnt – keine Abnahme! Fristbeginn gesondert klären")
    if re.search(r"fiktive\s+Abnahme|§\s*640\s+Abs\.?\s*2", volltext, re.IGNORECASE):
        pruef.append("Fiktive Abnahme § 640 Abs. 2 erwähnt – Fristbeginn rechtlich prüfen")

    # 3) Freitextfelder
    for feld, muster in FELD_MUSTER.items():
        wert, beleg = erstes_treffer_feld(volltext, muster)
        if wert:
            zeile[feld] = wert
            fund.append(f"{feld}: „{beleg}“")
        elif feld in ("Projekt", "Auftragnehmer"):
            pruef.append(f"{feld} nicht erkannt")

    # 4) Vertragsart
    va, beleg = vertragsart_bestimmen(volltext)
    zeile["Vertragsart"] = va
    if va:
        fund.append(f"Vertragsart: {beleg}")
        if va == "VOB/B":
            pruef.append("VOB/B: 2-Jahres-Frist für maschinelle/elektrotechnische Anlagen ohne Wartungsvertrag prüfen (§ 13 Abs. 4 Nr. 2 VOB/B)")
    else:
        pruef.append("Vertragsart nicht erkannt – entscheidet über 4 oder 5 Jahre!")

    # 5) Adresse / Ort / Region
    adr, plz, ort, beleg = adresse_finden(volltext)
    zeile["Adresse"], zeile["PLZ"], zeile["Ort"] = adr, plz, ort
    if beleg:
        fund.append(f"Adresse: „{beleg}“")
    zeile["Region"] = region_bestimmen(ort)
    if not zeile["Region"]:
        pruef.append(f"Region nicht als Rhein-Main erkannt (Ort: {ort or 'unbekannt'}) – sonst nicht im Radar")

    # 6) Rolle AN heuristisch, aber nur bei klarem Beleg
    if re.search(r"Generalunternehmer|\bGU\b|schlüsselfertig", volltext, re.IGNORECASE):
        zeile["Rolle AN"] = "GU"
    elif re.search(r"Generalübernehmer|\bGÜ\b", volltext, re.IGNORECASE):
        zeile["Rolle AN"] = "GÜ"

    zeile["Status"] = "offen"
    zeile["Fristende"] = ""   # bewusst leer: wird berechnet oder manuell gepflegt
    pruef.append("Fristende leer = wird berechnet; vertragliche Sonderfrist? dann hier eintragen")

    zeile["Bemerkung"] = f"Automatisch importiert aus {pfad.name} am {dt.date.today().isoformat()} – vor Nutzung prüfen"
    zeile["_Quelldatei"] = pfad.name
    zeile["_Fundstellen"] = " | ".join(fund)
    zeile["_Pruefbedarf"] = " | ".join(pruef)
    return zeile

# -------------------------------------------------------------------------- Lauf

def lauf(ordner, ziel):
    p = pathlib.Path(ordner)
    dateien = sorted(f for f in p.rglob("*") if f.suffix.lower() in (".pdf", ".docx", ".txt", ".md"))
    if not dateien:
        print(f"Keine auswertbaren Dokumente in {p}", file=sys.stderr)
        return 1
    zeilen = []
    for f in dateien:
        try:
            zeilen.append(dokument_auswerten(f))
            print(f"gelesen: {f.name}")
        except Exception as e:
            print(f"FEHLER  {f.name}: {type(e).__name__}: {e}", file=sys.stderr)
    with open(ziel, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=SPALTEN + ZUSATZ, delimiter=";")
        w.writeheader()
        w.writerows(zeilen)

    voll = [z for z in zeilen if z["Abnahmedatum"] and z["Vertragsart"] and z["Region"]]
    print(f"\n{len(zeilen)} Dokument(e) ausgewertet -> {ziel}")
    print(f"  {len(voll)} Zeile(n) fristenreif (Abnahme + Vertragsart + Region belegt)")
    print(f"  {len(zeilen)-len(voll)} Zeile(n) mit offenem Prüfbedarf\n")
    for z in zeilen:
        kennz = "OK " if z in voll else "PRÜF"
        print(f"  [{kennz}] {z['_Quelldatei']}: Abnahme={z['Abnahmedatum'] or '—'} "
              f"Vertrag={z['Vertragsart'] or '—'} Ort={z['Ort'] or '—'}")
        for hinweis in z["_Pruefbedarf"].split(" | "):
            if hinweis:
                print(f"         · {hinweis}")
    return 0

# ---------------------------------------------------------------------- Selbsttest

SELFTEST_DOKS = {
    "TEST_abnahmeprotokoll_vob.txt": """
TESTDOKUMENT – KEIN ECHTER VORGANG
Förmliches Abnahmeprotokoll

Bauvorhaben: Wohnquartier Testfeld, Bauabschnitt 2
Musterweg 14, 60329 Frankfurt am Main

Auftraggeber: Testeigentümer GmbH & Co. KG
Auftragnehmer: Testbau Generalunternehmung GmbH
Leistung: Schlüsselfertige Errichtung von 62 Wohneinheiten
Vertragsgrundlage: BGB-Bauvertrag unter Einbeziehung der VOB/B

Tag der Abnahme: 12.03.2024
Die Abnahme erfolgt unter Vorbehalt der Vertragsstrafe.
Sicherheit: Gewährleistungsbürgschaft 5 %, 310.000 EUR
""",
    "TEST_abnahmeprotokoll_bgb.txt": """
TESTDOKUMENT – KEIN ECHTER VORGANG
Abnahmeniederschrift

Objekt: Verwaltungsgebäude Testhausen
Prüfstraße 3, 65203 Wiesbaden

Bauherr: Test Invest AG
Auftragnehmer: Muster Ausbau GmbH
Gewerk: Innenausbau und Trockenbau
Der Vertrag ist ein Werkvertrag nach BGB.

Die Abnahme fand am 4. Juli 2023 statt.
Teilabnahme der TGA-Anlagen: abgenommen am 15.02.2023
""",
    "TEST_ohne_datum.txt": """
TESTDOKUMENT – KEIN ECHTER VORGANG
Schlussrechnung

Bauvorhaben: Halle Testort
Auftragnehmer: Testhallenbau GmbH
Ein Abnahmetermin ist in diesem Dokument nicht genannt.
""",
}

def selftest():
    ordner = pathlib.Path("_selftest_abnahmen")
    ordner.mkdir(exist_ok=True)
    for name, inhalt in SELFTEST_DOKS.items():
        (ordner / name).write_text(inhalt, encoding="utf-8")
    rc = lauf(ordner, "_selftest_rohimport.csv")
    print("\n--- Erwartungsabgleich ---")
    import csv as _csv
    with open("_selftest_rohimport.csv", encoding="utf-8") as fh:
        got = {r["_Quelldatei"]: r for r in _csv.DictReader(fh, delimiter=";")}
    erwartet = [
        ("TEST_abnahmeprotokoll_vob.txt", "Abnahmedatum", "2024-03-12"),
        ("TEST_abnahmeprotokoll_vob.txt", "Vertragsart", "VOB/B"),
        ("TEST_abnahmeprotokoll_vob.txt", "Region", "Rhein-Main"),
        ("TEST_abnahmeprotokoll_bgb.txt", "Abnahmedatum", "2023-07-04"),
        ("TEST_abnahmeprotokoll_bgb.txt", "Vertragsart", "BGB"),
        ("TEST_ohne_datum.txt", "Abnahmedatum", ""),
    ]
    fehler = 0
    for datei, feld, soll in erwartet:
        ist = got[datei][feld]
        ok = ist == soll
        fehler += 0 if ok else 1
        print(f"  {'✓' if ok else '✗'} {datei} · {feld}: erwartet „{soll}“, erhalten „{ist}“")
    print(f"\n{'Selbsttest bestanden.' if not fehler else f'{fehler} Abweichung(en)!'}")
    return rc or (1 if fehler else 0)

if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("ordner", nargs="?", help="Ordner mit Abnahmeprotokollen/Verträgen")
    ap.add_argument("-o", "--out", default="rohimport.csv")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    sys.exit(selftest() if a.selftest else (lauf(a.ordner, a.out) if a.ordner else ap.print_help() or 2))
